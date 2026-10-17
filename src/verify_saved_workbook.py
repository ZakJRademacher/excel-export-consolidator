"""Read-only comparison of a saved native Excel refresh with current input CSVs.

This does not refresh Excel. Run it only after Refresh All finishes and Save.
Requires openpyxl; the normal workbook and reference tests do not.
"""
import argparse
from collections import Counter
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import zipfile

import openpyxl
from consolidate import ROOT, FIELDS, consolidate


def verify(stage):
    path = ROOT / "deliverables/Export Consolidator.xlsx"
    reference = consolidate(ROOT / "data/incoming")
    wb = openpyxl.load_workbook(path, data_only=True)

    def table(sheet, name):
        ws = wb[sheet]
        cells = list(ws[ws.tables[name].ref])
        headers = [c.value for c in cells[0]]
        return [dict(zip(headers, [c.value for c in row])) for row in cells[1:]
                if any(c.value is not None for c in row)]

    clean = table("Clean master", "CleanMaster")
    review = table("Needs review", "NeedsReview")
    history = table("Import history", "ImportHistory")
    expected = [r for r in reference["rows"] if r["Status"] == "Clean"]
    assert len(clean) == len(expected)
    by_id = {str(r["RowID"]): r for r in clean}
    assert len(by_id) == len(clean), "Duplicate accepted RowID"
    for row in expected:
        actual = by_id[row["RowID"]]
        for field in FIELDS + ["SourceFile", "SourceRow"]:
            value = actual[field]
            if field in ("OrderDate", "ShipDate"):
                value = value.strftime("%Y-%m-%d")
            elif field == "Sales":
                value = Decimal(str(value)).quantize(Decimal("0.0001"))
                assert value == Decimal(row[field]), (row["RowID"], field, value)
                continue
            assert str(value) == str(row[field]), (row["RowID"], field, value)
    actual_keys = Counter((r["Status"], r["SourceFile"], int(r["SourceRow"]), str(r["RowID"] or "")) for r in review)
    expected_keys = Counter((r["Status"], r["SourceFile"], r["SourceRow"], str(r["RowID"] or "")) for r in reference["rows"] if r["Status"] != "Clean")
    assert actual_keys == expected_keys, "Review/repeat lineage differs"
    for native in history:
        match = next(f for f in reference["files"] if f["SourceFile"] == native["SourceFile"])
        for key in ["InputRows", "CleanRows", "ReviewRows", "RepeatedRows", "FileIssues"]:
            assert native[key] == match[key], (native["SourceFile"], key)
    assert len(history) == len(reference["files"])
    metrics = reference["metrics"]
    for cell, key in {"C8":"files", "C9":"input_rows", "C10":"clean_rows", "C11":"review_rows", "C12":"repeated_rows", "C13":"file_issues"}.items():
        assert wb["Summary"][cell].value == metrics[key], (cell, key)
    total = sum((Decimal(str(r["Sales"])).quantize(Decimal("0.0001")) for r in clean), Decimal(0))
    assert total == Decimal(metrics["clean_sales"])
    assert Decimal(str(wb["Summary"]["C16"].value)).quantize(Decimal("0.0001")) == total
    errors = [(ws.title, c.coordinate, c.value) for ws in wb for row in ws for c in row if c.data_type == "e"]
    assert not errors, errors
    with zipfile.ZipFile(path) as z:
        assert "xl/connections.xml" in z.namelist()
        assert len([n for n in z.namelist() if n.startswith("xl/queryTables/queryTable") and n.endswith(".xml")]) == 3
    result = {"stage": stage, "metrics": metrics, "accepted_fields_and_lineage_match": True,
              "review_status_and_lineage_match": True, "per_file_counts_match": True,
              "summary_cached_values_match": True, "formula_errors": errors,
              "worksheet_query_tables": 3, "workbook_sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
    (ROOT / "evidence" / ("native-" + stage + ".json")).write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps(result))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("stage", choices=["baseline", "added-file", "unchanged-repeat", "file-issue", "restored"])
    verify(parser.parse_args().stage)
